#!/usr/bin/env python3
"""One-pass, metadata/header-only Dryad source gate for an external BAM candidate.

This script never scores butterfly or host-plant outcomes. Failed file acquisition
is a terminal STOP, not a request for automatic retries or source substitution.
"""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from tempfile import TemporaryDirectory
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
import zipfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PROTOCOL = ROOT / "validation/gotland_bam_preflight_v1/protocol_v1.json"
PROTOCOL_COMMIT = "7d77af98d6c0ecf7fcdc9f99897c173bfb219028"
MIN_BYTES, MAX_BYTES = 4096, 5242880


def _digest(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def audit_workbook_header(payload: bytes, filename: str) -> dict[str, Any]:
    """Read workbook structure plus first worksheet row only; never scientific data."""
    if not (MIN_BYTES <= len(payload) <= MAX_BYTES):
        raise ValueError("SOURCE_STOP_SIZE_OUTSIDE_DECLARED_RANGE")
    if not payload.startswith(b"PK") or not zipfile.is_zipfile(__import__("io").BytesIO(payload)):
        raise ValueError("SOURCE_STOP_INVALID_XLSX_ZIP")
    import io
    stream = io.BytesIO(payload)
    with zipfile.ZipFile(stream) as z:
        files = set(z.namelist())
        if "xl/workbook.xml" not in files or "[Content_Types].xml" not in files:
            raise ValueError("SOURCE_STOP_NOT_XLSX_WORKBOOK")
        if any(n.endswith("vbaProject.bin") for n in files):
            raise ValueError("SOURCE_STOP_UNEXPECTED_MACRO")
    stream.seek(0)
    book = load_workbook(stream, read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in book.worksheets:
            # Physical first row ONLY; never inspect rows 2+.
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = ["" if value is None else str(value).strip() for value in first]
            if not any(headers):
                raise ValueError("SOURCE_STOP_EMPTY_FIRST_PHYSICAL_HEADER")
            sheets.append({
                "sheet": sheet.title,
                "column_count": len(headers),
                "row_count_from_workbook_dimensions": sheet.max_row,
                "physical_headers": headers,
            })
        if not sheets:
            raise ValueError("SOURCE_STOP_WORKBOOK_WITHOUT_SHEETS")
    finally:
        book.close()
    return {
        "source_file_name": filename,
        "sha256": _digest(payload),
        "size_bytes": len(payload),
        "sheets": sheets,
        "data_rows_examined": 0,
        "biological_response_computed": False,
    }


def _fetch_once(url: str) -> bytes:
    request = Request(
        url,
        headers={"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    with urlopen(request, timeout=35) as response:
        status = response.status
        if status != 200:
            raise ValueError(f"SOURCE_STOP_HTTP_{status}")
        content = response.read(MAX_BYTES + 1)
    return content


def schema_crosswalk_status(workbooks: list[dict[str, Any]]) -> dict[str, Any]:
    if len(workbooks) != 2:
        raise ValueError("expected exactly the two frozen workbooks")
    selected = []
    for f in workbooks:
        heads = {
            normalize_header(h)
            for s in f["sheets"] for h in s["physical_headers"]
            if h
        }
        selected.append(heads)
    common = sorted(selected[0] & selected[1])
    candidates = [
        value for value in common
        if any(tag in value for tag in ("patch", "site", "parcel", "grid", "loc"))
    ]
    return {
        "identical_header_tokens_across_workbooks": common,
        "candidate_common_location_tokens": candidates,
        "host_to_patch_crosswalk_verified": False,
        "B_status": "B_UNQUALIFIED_PENDING_INDEPENDENT_PATCH_CROSSWALK",
        "explanation": (
            "Shared header labels are not proof that plot-grid locations and "
            "256 butterfly patch locations refer to the same unit. "
            "No biological row join is authorized."
        ),
    }


def run_gate(*, protocol_path: Path = PROTOCOL, fetch=_fetch_once) -> dict[str, Any]:
    raw_protocol = protocol_path.read_bytes()
    protocol = json.loads(raw_protocol)
    if protocol.get("schema") != "eog.gotland_bam_pre_response_source_gate.v1":
        raise ValueError("unrecognized frozen source protocol")
    entries = protocol["source"]["workbooks"]
    if [(item["file_id"], item["file_name"]) for item in entries] != [
        (264459, "Butterfly_occurrence_data.xlsx"),
        (264460, "Host_plant_data.xlsx"),
    ]:
        raise ValueError("source identity altered; refusing acquisition")
    receipt = {
        "schema": "eog.gotland_bam_schema_receipt.v1",
        "source": protocol["source"]["dataset_doi"],
        "protocol_commit": PROTOCOL_COMMIT,
        "protocol_sha256": _digest(raw_protocol),
        "status": "SCHEMA_GATE_STARTED",
        "response_rows_opened": 0,
        "scored_ecological_endpoints": 0,
        "file_results": [],
    }
    for entry in entries:
        try:
            payload = fetch(entry["url"])
            info = audit_workbook_header(payload, entry["file_name"])
            info["source_file_id"] = entry["file_id"]
            receipt["file_results"].append(info)
        except (HTTPError, URLError, OSError, ValueError, zipfile.BadZipFile) as exc:
            if isinstance(exc, HTTPError):
                status = f"SOURCE_STOP_HTTP_{exc.code}"
            elif isinstance(exc, URLError):
                status = "SOURCE_STOP_TRANSPORT"
            else:
                status = str(exc) if str(exc).startswith("SOURCE_STOP_") else "SOURCE_STOP_SCHEMA_OR_TRANSPORT"
            receipt["status"] = status
            receipt["stopped_at_file_id"] = entry["file_id"]
            receipt["crosswalk"] = None
            receipt["fail_closed_no_fallback"] = True
            receipt["fingerprint"] = _digest(json.dumps(receipt, sort_keys=True).encode())
            return receipt
    receipt["status"] = "HEADER_ONLY_SUCCESS__BIOLOGICAL_JOIN_NOT_YET_QUALIFIED"
    receipt["crosswalk"] = schema_crosswalk_status(receipt["file_results"])
    receipt["fingerprint"] = _digest(json.dumps(receipt, sort_keys=True).encode())
    return receipt


def main() -> None:
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--receipt", required=True, type=Path)
    args = parser.parse_args()
    result = run_gate()
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print("GOTLAND_SOURCE_GATE="+json.dumps({
        "status": result["status"],
        "source": result["source"],
        "n_workbooks_verified": len(result["file_results"]),
        "response_rows_opened": result["response_rows_opened"],
        "fingerprint": result["fingerprint"],
    }, sort_keys=True))


if __name__ == "__main__":
    main()
